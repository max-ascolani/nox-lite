#!/usr/bin/env python3
"""Create a clean DOCX from reviewed CV JSON; optionally convert with LibreOffice."""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from urllib.parse import urlsplit


def validate_cv(data):
    if not isinstance(data, dict) or data.get("schema_version") != 1:
        raise ValueError("CV must be a JSON object with schema_version 1")
    if not isinstance(data.get("name"), str) or not data["name"].strip():
        raise ValueError("A confirmed name is required")
    if "{{" in json.dumps(data):
        raise ValueError("Replace all template placeholders before building")
    for field in ("location", "email", "phone", "summary"):
        if not isinstance(data.get(field, ""), str):
            raise ValueError(f"{field} must be text")
    if not data.get("experience") and not data.get("education"):
        raise ValueError("Include relevant experience or education")
    for key in ("experience", "education", "skills", "links"):
        if not isinstance(data.get(key, []), list):
            raise ValueError(f"{key} must be a list")
    for item in data.get("experience", []):
        if not isinstance(item, dict):
            raise ValueError("Each experience must be an object")
        for field in ("organisation", "role", "dates"):
            if not isinstance(item.get(field), str) or not item[field].strip():
                raise ValueError(f"Experience requires {field}; use an honest year-only range if months are unresolved")
        if not isinstance(item.get("bullets", []), list) or not all(isinstance(x, str) and x.strip() for x in item.get("bullets", [])):
            raise ValueError("Experience bullets must be nonempty strings")
        if not isinstance(item.get("location", ""), str) or not isinstance(item.get("description", ""), str):
            raise ValueError("Experience location and description must be text")
    for item in data.get("education", []):
        if not isinstance(item, dict):
            raise ValueError("Each education entry must be an object")
        if not all(isinstance(item.get(k), str) and item[k].strip() for k in ("institution", "qualification", "dates")):
            raise ValueError("Education requires institution, qualification and dates")
        if not isinstance(item.get("detail", ""), str):
            raise ValueError("Education detail must be text")
    for item in data.get("skills", []):
        if not isinstance(item, dict):
            raise ValueError("Each skill group must be an object")
        if not isinstance(item.get("category"), str) or not isinstance(item.get("items"), list) or not all(isinstance(x, str) for x in item["items"]):
            raise ValueError("Skills need a category and text items")
    for link in data.get("links", []):
        if not isinstance(link, dict) or not isinstance(link.get("url"), str) or not isinstance(link.get("label"), str):
            raise ValueError("Links need text label and url fields")
        parsed = urlsplit(link.get("url", ""))
        if parsed.scheme not in {"https", "http"} or not parsed.hostname or parsed.username or parsed.password or not link.get("label"):
            raise ValueError("Profile links need a label and public http(s) URL without credentials")


def find_soffice(executable=None):
    """Return LibreOffice from an explicit path, PATH, or a standard macOS/Windows install; else None."""
    if executable:
        return shutil.which(executable)
    candidates = ["/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    for variable, default in (("ProgramFiles", r"C:\Program Files"), ("ProgramFiles(x86)", r"C:\Program Files (x86)")):
        candidates.append(os.path.join(os.environ.get(variable, default), "LibreOffice", "program", "soffice.exe"))
    return shutil.which("soffice") or next((c for c in candidates if os.path.isfile(c) and os.access(c, os.X_OK)), None)


def pin_font(styles, names, font):
    """Theme attributes (w:asciiTheme etc.) outrank w:ascii, so remove them and name the font for every script."""
    from docx.oxml.ns import qn
    for fonts in list(styles.element.iter(qn("w:rFonts"))):
        for key in [k for k in fonts.attrib if k.lower().endswith("theme")]:
            del fonts.attrib[key]
        if not fonts.attrib:
            fonts.getparent().remove(fonts)
    defaults = styles.element.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
    for rpr in [defaults] + [styles[name].element.get_or_add_rPr() for name in names]:
        if rpr is not None:
            fonts = rpr.get_or_add_rFonts()
            for script in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn("w:" + script), font)


FONT = "Georgia"  # A classic serif installed on every Mac and Windows computer.
LINK_COLOUR = "0563C1"
# w:pPr children that must follow w:pBdr, in schema order (Word rejects out-of-order XML).
AFTER_BORDER = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
                "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
                "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def new_document(paper, font, sizes, margins):
    """A blank document with the stock template's extras removed and the chosen font everywhere."""
    if paper not in {"a4", "letter"} or not isinstance(font, str) or not font.strip():
        raise ValueError("Choose a4 or letter paper and a nonempty font name")
    try:
        from docx import Document
        from docx.shared import Inches, Mm, Pt, RGBColor
        from docx.oxml.ns import qn
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
    except ImportError as exc:
        raise ValueError("Install requirements-docs.txt in an isolated Python environment, or use your agent's bundled document runtime.") from exc
    doc = Document()
    # Remove irrelevant stock template parts, including an out-of-date preview image.
    for rel_id, rel in list(doc.part.rels.items()):
        if rel.reltype == RT.CUSTOM_XML:
            doc.part.drop_rel(rel_id)
    for rel_id, rel in list(doc.part.package.rels.items()):
        if rel.reltype == RT.THUMBNAIL:
            del doc.part.package.rels[rel_id]
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = (Mm(210), Mm(297)) if paper == "a4" else (Inches(8.5), Inches(11))
    section.top_margin = section.bottom_margin = margins[0]
    section.left_margin = section.right_margin = margins[1]
    pin_font(doc.styles, list(sizes), font)
    for name, size in sizes.items():
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.widow_control = True
    return doc


def border(paragraph, *sides, val="single", size=6, space=1, color="000000"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    edges = OxmlElement("w:pBdr")
    for side in sides:
        edge = OxmlElement("w:" + side)
        for key, value in (("val", val), ("sz", size), ("space", space), ("color", color)):
            edge.set(qn("w:" + key), str(value))
        edges.append(edge)
    paragraph._p.get_or_add_pPr().insert_element_before(edges, *AFTER_BORDER)


def add_link(paragraph, label, url, size):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    hyper = OxmlElement("w:hyperlink")
    hyper.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run, props = OxmlElement("w:r"), OxmlElement("w:rPr")
    for tag, value in (("w:color", LINK_COLOUR), ("w:sz", str(size * 2)), ("w:szCs", str(size * 2)), ("w:u", "single")):
        item = OxmlElement(tag)
        item.set(qn("w:val"), value)
        props.append(item)
    text = OxmlElement("w:t")
    text.text = label
    run.append(props)
    run.append(text)
    hyper.append(run)
    paragraph._p.append(hyper)


def header(doc, data, size=10):
    """Centred name, then one contact line between dashed rules. The cover letter uses it too."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    section = doc.sections[0]
    name = doc.add_paragraph(data["name"], "Title")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.line_spacing = 1.0
    name.paragraph_format.keep_with_next = True
    parts = [data.get(k, "").strip() for k in ("location", "phone") if data.get(k, "").strip()]
    if data.get("email", "").strip():
        parts.append((data["email"].strip(), "mailto:" + data["email"].strip()))
    parts += [(link["label"], link["url"]) for link in data.get("links", [])]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    # Keep the rules about as wide as the text: roughly half an em per character, plus some margin.
    shown = " | ".join(part[0] if isinstance(part, tuple) else part for part in parts)
    width = section.page_width - section.left_margin - section.right_margin
    spare = width - Pt(size * 0.6 * len(shown))
    if spare > 0:
        p.paragraph_format.left_indent = p.paragraph_format.right_indent = int(spare / 2)
    border(p, "top", "bottom", val="dashSmallGap", size=4, space=4)
    for index, part in enumerate(parts):
        if index:
            p.add_run(" | ").font.size = Pt(size)
        if isinstance(part, tuple):
            add_link(p, part[0], part[1], size)
        else:
            p.add_run(part).font.size = Pt(size)
    return p


def build(data, output, paper="a4", font=FONT, replace=False):
    validate_cv(data)
    output = Path(output).resolve()
    if output.suffix.lower() != ".docx":
        raise ValueError("Output must end in .docx")
    if output.exists() and not replace:
        raise ValueError("Output exists. Reconcile any direct Word edits first; use --replace only intentionally.")
    sizes = {"Normal": 10.5, "Title": 18, "Heading 1": 10, "Heading 2": 10.5, "List Bullet": 10}
    try:
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
        from docx.shared import Cm, Pt
    except ImportError as exc:
        raise ValueError("Install requirements-docs.txt in an isolated Python environment, or use your agent's bundled document runtime.") from exc
    doc = new_document(paper, font, sizes, (Cm(0.9), Cm(1.0)))
    section = doc.sections[0]
    styles = doc.styles
    styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].font.bold = name != "Heading 2"
        styles[name].font.italic = False
        styles[name].paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    width = section.page_width - section.left_margin - section.right_margin
    styles["Heading 2"].paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    styles["List Bullet"].paragraph_format.left_indent = Cm(1.0)
    styles["List Bullet"].paragraph_format.first_line_indent = Cm(-0.5)

    def heading(text):
        border(doc.add_paragraph(text.upper(), "Heading 1"), "bottom")

    def dated(name, place, dates):
        p = doc.add_paragraph(style="Heading 2")
        p.add_run(name).bold = True
        if place:
            p.add_run(", " + place)
        p.add_run("\t" + dates)

    def line(text, keep, bold=False, italic=False, size=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        p.paragraph_format.keep_with_next = keep
        return p

    header(doc, data)
    if data.get("summary", "").strip():
        heading("Profile")
        doc.add_paragraph(data["summary"].strip())
    if data.get("experience"):
        heading("Experience")
        for item in data["experience"]:
            bullets = item.get("bullets", [])
            description = item.get("description", "").strip()
            dated(item["organisation"], item.get("location", "").strip(), item["dates"])
            line(item["role"], bool(bullets or description), bold=True)
            if description:
                line(description, bool(bullets), italic=True)
            for bullet in bullets:
                doc.add_paragraph(bullet, "List Bullet").paragraph_format.keep_together = True
    if data.get("education"):
        heading("Education")
        for item in data["education"]:
            dated(item["institution"], "", item["dates"])
            line(item["qualification"], bool(item.get("detail")))
            if item.get("detail"):
                line(item["detail"], False)
    if data.get("skills"):
        heading("Skills")
        for item in data["skills"]:
            p = doc.add_paragraph()
            p.add_run(item["category"] + ": ").bold = True
            p.add_run(", ".join(item["items"]))
            for run in p.runs:
                run.font.size = Pt(10)
    cp = doc.core_properties
    cp.title = data["name"] + " CV"
    cp.author = data["name"]
    cp.last_modified_by = ""
    cp.subject = cp.keywords = cp.comments = cp.category = ""
    cp.created = cp.modified = datetime.now(timezone.utc)
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, temp = tempfile.mkstemp(suffix=".docx", dir=output.parent)
    os.close(fd)
    try:
        doc.save(temp)
        os.replace(temp, output)
    finally:
        if os.path.exists(temp): os.unlink(temp)
    return output


def convert_pdf(docx, executable, replace=False):
    output = Path(docx).with_suffix(".pdf")
    if output.exists() and not replace:
        raise ValueError("PDF exists; use --replace only after reconciling edits")
    found = find_soffice(executable)
    if not found:
        raise ValueError("DOCX created. PDF needs LibreOffice: provide --soffice with its executable path, or export from Word and inspect every page.")
    with tempfile.TemporaryDirectory(prefix="nox-render-") as directory:
        root = Path(directory)
        profile = (root / "profile").as_uri()
        result = subprocess.run([found, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf", "--outdir", str(root), str(docx)], capture_output=True, text=True, timeout=120)
        generated = root / output.name
        if result.returncode or not generated.exists() or generated.stat().st_size == 0:
            raise ValueError("PDF conversion failed; DOCX retained. " + result.stderr[-600:])
        shutil.copy2(generated, output)
    return output


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("source"); p.add_argument("output")
    p.add_argument("--paper", choices=["a4", "letter"], default="a4")
    p.add_argument("--font", default=FONT); p.add_argument("--replace", action="store_true")
    p.add_argument("--pdf", action="store_true"); p.add_argument("--soffice")
    args = p.parse_args()
    try:
        data = json.loads(Path(args.source).read_text(encoding="utf-8"))
        output = build(data, args.output, args.paper, args.font, args.replace)
        print(f"Created {output}")
        if args.pdf: print(f"Created {convert_pdf(output, args.soffice, args.replace)}")
        print("Render and inspect every page; generation alone does not establish layout quality or human approval.")
    except (ValueError, OSError, json.JSONDecodeError, subprocess.TimeoutExpired) as exc:
        print(f"Nox-Lite: {exc}", file=sys.stderr); raise SystemExit(1)


if __name__ == "__main__": main()
