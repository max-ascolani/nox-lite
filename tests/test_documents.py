import copy
import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import build_cv
import nox


class ValidationTests(unittest.TestCase):
    def setUp(self):
        self.data = nox.read_json(nox.ROOT / "examples/candidate/cv.json")

    def test_optional_contacts_and_education_only(self):
        self.data.update(email="", phone="", links=[], experience=[])
        build_cv.validate_cv(self.data)

    def test_invalid_nested_data_has_useful_errors(self):
        for field, bad in [("experience", ["not an object"]), ("education", [None]), ("skills", [42]), ("links", ["url"]), ("links", [{"url": 123, "label": "Work"}])]:
            data = copy.deepcopy(self.data)
            data[field] = bad
            with self.assertRaises(ValueError):
                build_cv.validate_cv(data)

    def test_rejects_unfinished_or_invalid_document(self):
        for field, value in [("name", ""), ("summary", "{{not completed}}"), ("email", 123), ("experience", None)]:
            data = copy.deepcopy(self.data)
            data[field] = value
            with self.assertRaises(ValueError):
                build_cv.validate_cv(data)
        self.data["links"] = [{"label": "private", "url": "https://user:secret@example.com"}]
        with self.assertRaises(ValueError):
            build_cv.validate_cv(self.data)


@unittest.skipUnless(importlib.util.find_spec("docx"), "Install requirements-docs.txt for document tests")
class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="nox-cv-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.data = nox.read_json(nox.ROOT / "examples/candidate/cv.json")

    def test_a4_letter_unicode_and_contact_are_portable(self):
        from docx import Document
        self.data["name"] = "Jordán Rivera"
        for paper, expected_width in [("a4", 210), ("letter", 215.9)]:
            path = build_cv.build(self.data, self.root / (paper + ".docx"), paper=paper)
            doc = Document(path)
            self.assertAlmostEqual(doc.sections[0].page_width.mm, expected_width, places=1)
            self.assertEqual(doc.paragraphs[0].text, "Jordán Rivera")
            self.assertEqual(doc.core_properties.author, "Jordán Rivera")
            self.assertEqual(doc.core_properties.last_modified_by, "")
            self.assertIn("jordan@example.com", "\n".join(p.text for p in doc.paragraphs))
            with zipfile.ZipFile(path) as archive:
                rels = archive.read("word/_rels/document.xml.rels").decode()
                self.assertIn("https://example.com/jordan", rels)
                self.assertFalse(any(name.startswith("word/embeddings/") for name in archive.namelist()))
                self.assertFalse(any(name.startswith("customXml/") for name in archive.namelist()))
                self.assertNotIn("docProps/thumbnail.jpeg", archive.namelist())
                self.assertNotIn(b"<w:pBdr>", archive.read("word/styles.xml"))

    def test_name_and_headings_use_the_chosen_font(self):
        import xml.etree.ElementTree as ET
        w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        for font in ("Arial", "Liberation Sans"):
            path = build_cv.build(self.data, self.root / (font + ".docx"), font=font)
            with zipfile.ZipFile(path) as archive:
                raw = archive.read("word/styles.xml")
            # Theme attributes outrank w:ascii, which made the name and headings render in Calibri/Carlito.
            self.assertNotIn(b"asciiTheme", raw)
            self.assertNotIn(b"hAnsiTheme", raw)
            styles = ET.fromstring(raw)
            by_id = {s.get(w + "styleId"): s for s in styles.iter(w + "style")}
            targets = [styles.find(f"{w}docDefaults/{w}rPrDefault/{w}rPr")] + [by_id[i] for i in ("Normal", "Title", "Heading1", "Heading2", "ListBullet")]
            for target in targets:
                fonts = target.find(f".//{w}rFonts")
                self.assertEqual({fonts.get(w + k) for k in ("ascii", "hAnsi", "eastAsia", "cs")}, {font})

    def test_direct_word_changes_cannot_be_overwritten_accidentally(self):
        path = build_cv.build(self.data, self.root / "cv.docx")
        before = path.read_bytes()
        with self.assertRaisesRegex(ValueError, "Reconcile"):
            build_cv.build(self.data, path)
        self.assertEqual(path.read_bytes(), before)

    def test_missing_pdf_converter_preserves_docx(self):
        path = build_cv.build(self.data, self.root / "cv.docx")
        with self.assertRaisesRegex(ValueError, "DOCX created"):
            build_cv.convert_pdf(path, "nonexistent-nox-test-converter")
        self.assertTrue(path.exists())
        self.assertFalse(path.with_suffix(".pdf").exists())


if __name__ == "__main__":
    unittest.main()
